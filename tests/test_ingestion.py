"""
Automated Tests for Document Ingestion & Text Extraction Layer
Covers TXT, DOCX, Digital PDF, scanned detection, corrupted documents, and file constraints.
"""

import io
import pytest
import docx
import fitz

from core.ingestion.parser import document_parser, DocumentParsingError, ExtractionResult
from core.ingestion.ocr import OCREngine


def test_txt_extraction_valid():
    sample_text = "Experienced Python Developer with expertise in Django, FastAPI, and PostgreSQL."
    txt_bytes = sample_text.encode("utf-8")
    result = document_parser.parse_file(txt_bytes, "resume.txt")

    assert result.success is True
    assert result.file_format == "TXT"
    assert "Python Developer" in result.raw_text
    assert result.word_count > 5
    assert result.character_count > 20


def test_docx_extraction_valid():
    doc = docx.Document()
    doc.add_heading("Senior Data Scientist Resume", level=1)
    doc.add_paragraph("Machine learning specialist with deep experience in NLP and PyTorch.")
    
    # Add a table
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Skills"
    hdr_cells[1].text = "Python, Scikit-Learn, Pandas"

    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_bytes = docx_buffer.getvalue()

    result = document_parser.parse_file(docx_bytes, "candidate.docx")
    assert result.success is True
    assert result.file_format == "DOCX"
    assert "Senior Data Scientist" in result.raw_text
    assert "PyTorch" in result.raw_text
    assert "Scikit-Learn" in result.raw_text


def test_pdf_digital_extraction_valid():
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "DevOps Engineer with Kubernetes, Docker, and AWS Cloud experience.")
    pdf_bytes = doc.tobytes()
    doc.close()

    result = document_parser.parse_file(pdf_bytes, "resume.pdf")
    assert result.success is True
    assert result.file_format == "PDF"
    assert "DevOps Engineer" in result.raw_text
    assert "Kubernetes" in result.raw_text
    assert result.ocr_used is False


def test_empty_file_rejection():
    with pytest.raises(DocumentParsingError) as exc_info:
        document_parser.parse_file(b"", "empty.txt")
    assert "empty" in str(exc_info.value).lower()


def test_unsupported_file_type_rejection():
    with pytest.raises(DocumentParsingError) as exc_info:
        document_parser.parse_file(b"some content", "image.png")
    assert "unsupported format" in str(exc_info.value).lower()


def test_corrupted_pdf_handling():
    corrupted_bytes = b"%PDF-1.4 ... totally corrupted binary junk that cannot parse"
    with pytest.raises(DocumentParsingError):
        document_parser.parse_file(corrupted_bytes, "corrupted.pdf")


def test_ocr_availability_diagnostic():
    ocr = OCREngine()
    assert ocr.is_available is True
    status = ocr.get_status_info()
    assert status["available"] is True
    assert "engine" in status
    assert "RapidOCR" in status["engine"] or "Tesseract" in status["engine"]
